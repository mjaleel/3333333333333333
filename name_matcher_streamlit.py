# -*- coding: utf-8 -*-
"""
تطبيق مطابقة الرواتب — نسخة Streamlit
========================================
✓ واجهة ويب كاملة بدون Tkinter
✓ اختيار الأعمدة ديناميكي
✓ إعدادات مرنة (عتبة، وزن المدرسة)
✓ أعمدة إضافية من قاعدة البيانات
✓ نسبة المطابقة % مع شريط تقدم
✓ جدول نتائج ملوّن
✓ تصدير Excel احترافي (BytesIO)

تشغيل:
    streamlit run name_matcher_streamlit.py
"""

import re
import io
from datetime import datetime

import pandas as pd
import streamlit as st
from rapidfuzz import fuzz
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

# ══════════════════════════════════════════════════
#  إعداد الصفحة  (يجب أن يكون أول أمر Streamlit)
# ══════════════════════════════════════════════════

st.set_page_config(
    page_title="مطابقة الرواتب",
    page_icon="⚡",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ══════════════════════════════════════════════════
#  CSS مخصص — ثيم داكن عربي RTL
# ══════════════════════════════════════════════════

st.markdown("""
<style>
/* ── خط + اتجاه ── */
@import url('https://fonts.googleapis.com/css2?family=Tajawal:wght@400;500;700&display=swap');

html, body, [class*="css"] {
    font-family: 'Tajawal', sans-serif !important;
    direction: rtl;
}

/* ── خلفية عامة ── */
.stApp { background-color: #0d1117; }

/* ── الشريط الجانبي ── */
section[data-testid="stSidebar"] {
    background-color: #161b22 !important;
    border-left: 1px solid #30363d;
}
section[data-testid="stSidebar"] * { direction: rtl; }

/* ── بطاقات الإحصاء ── */
.stat-card {
    background: #21262d;
    border: 1px solid #30363d;
    border-radius: 10px;
    padding: 14px 10px;
    text-align: center;
}
.stat-num  { font-size: 2rem; font-weight: 700; margin: 4px 0 0; }
.stat-lbl  { font-size: .82rem; color: #8b949e; }

/* ── رأس الصفحة ── */
.page-header {
    background: linear-gradient(135deg, #1f2937 0%, #161b22 100%);
    border: 1px solid #30363d;
    border-radius: 12px;
    padding: 18px 24px;
    margin-bottom: 18px;
    display: flex;
    align-items: center;
    gap: 14px;
}
.page-header h1 {
    color: #e6edf3;
    font-size: 1.7rem;
    margin: 0;
}
.page-header span { font-size: 2rem; }

/* ── شريط التقدم ── */
.stProgress > div > div { background-color: #238636 !important; }

/* ── الجدول ── */
.dataframe-container { border-radius: 8px; overflow: hidden; }

/* ── أزرار ── */
.stButton > button {
    width: 100%;
    border-radius: 8px;
    font-family: 'Tajawal', sans-serif;
    font-weight: 700;
    font-size: 1rem;
    padding: 10px;
    transition: all .2s;
}

/* ── divider ── */
hr { border-color: #30363d; }

/* ── selectbox / multiselect labels ── */
label { color: #c9d1d9 !important; font-weight: 500; }

/* ── expander ── */
details summary {
    color: #8b949e;
    font-size: .9rem;
}

/* ── info boxes ── */
.stAlert { border-radius: 8px; }

/* ── تلوين صفوف الجدول ── */
.row-matched { background-color: #0d2818 !important; }
.row-missing { background-color: #2d0f0f !important; }
</style>
""", unsafe_allow_html=True)


# ══════════════════════════════════════════════════
#  منطق المطابقة
# ══════════════════════════════════════════════════

def normalize_arabic(text: str) -> str:
    if not text or isinstance(text, float):
        return ""
    text = str(text).strip()
    text = re.sub(r'[إأآا]', 'ا', text)
    text = re.sub(r'ة',      'ه', text)
    text = re.sub(r'ى',      'ي', text)
    text = re.sub(r'[ؤئ]',  'ي', text)
    text = re.sub(r'[\u064B-\u0652\u0670]', '', text)
    text = " ".join(text.split())
    text = re.sub(r'(عبد)([^\s])', r'\1 \2', text)
    return text.lower()


def name_match_score(name_req: str, name_db: str) -> float:
    r = normalize_arabic(name_req).split()
    d = normalize_arabic(name_db).split()
    if not r or not d:
        return 0.0
    if r[0] != d[0]:
        return 0.0

    base_score = fuzz.ratio(" ".join(r), " ".join(d))

    if len(r) >= 2 and len(d) >= 2:
        if r[1] != d[1]:
            return 0.0
        base_score += 5

    if len(r) >= 3 and len(d) >= 3:
        if r[2] != d[2]:
            return 0.0
        base_score += 5

    if len(r) >= 4 and len(d) >= 4:
        if r[3] != d[3]:
            return 0.0
        base_score += 5
    elif len(r) == 3 and len(d) > 3:
        if r == d[:3]:
            base_score += 10
    elif len(r) == 4 and len(d) == 3:
        if r[:3] == d:
            base_score += 8

    return min(100.0, base_score)


def build_index(records: list, name_col: str) -> dict:
    idx: dict = {}
    for rec in records:
        parts = normalize_arabic(rec.get(name_col, "")).split()
        p1  = parts[0] if len(parts) > 0 else "__"
        p3  = parts[2] if len(parts) > 2 else "__"
        key = (p1, p3)
        idx.setdefault(key, []).append(rec)
    return idx


def find_best_match(
    name_req: str, school_req: str,
    db_index: dict,
    db_name_col: str, db_school_col: str,
    use_school: bool, threshold: float, school_weight: float
) -> tuple:
    """يعيد (best_record, final_score, name_score, match_pct)"""
    parts = normalize_arabic(name_req).split()
    p1 = parts[0] if parts else "__"
    p3 = parts[2] if len(parts) > 2 else "__"

    candidates = db_index.get((p1, p3), [])
    if not candidates:
        for k, v in db_index.items():
            if k[0] == p1:
                candidates.extend(v)

    school_words: list = []
    if use_school and school_req:
        s_norm = normalize_arabic(school_req)
        school_words = [w for w in s_norm.split() if len(w) > 2]

    best_rec   = None
    best_final = 0.0
    best_name  = 0.0

    for rec in candidates:
        n_score = name_match_score(name_req, rec.get(db_name_col, ""))
        if n_score == 0:
            continue
        final = n_score
        if use_school and school_words:
            db_sn = normalize_arabic(rec.get(db_school_col, ""))
            hits  = sum(1 for w in school_words if w in db_sn)
            final += (hits / len(school_words)) * school_weight
        if final > best_final:
            best_final = final
            best_rec   = rec
            best_name  = n_score

    if best_final < threshold:
        return None, 0.0, 0.0, 0.0

    raw_pct = fuzz.ratio(
        normalize_arabic(name_req),
        normalize_arabic(best_rec.get(db_name_col, "") if best_rec else ""),
    )
    return best_rec, best_final, best_name, float(raw_pct)


# ══════════════════════════════════════════════════
#  توليد Excel بـ BytesIO
# ══════════════════════════════════════════════════

def build_excel(df: pd.DataFrame) -> bytes:
    buf = io.BytesIO()
    df.to_excel(buf, index=False, engine="openpyxl")
    buf.seek(0)

    wb = load_workbook(buf)
    ws = wb.active

    green_fill  = PatternFill("solid", fgColor="C6EFCE")
    red_fill    = PatternFill("solid", fgColor="FFC7CE")
    orange_fill = PatternFill("solid", fgColor="FFEB9C")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    bold_white  = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    normal_font = Font(name="Arial", size=10)

    # رأس الجدول
    for cell in ws[1]:
        cell.fill      = header_fill
        cell.font      = bold_white
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    # أعمدة الحالة والنسبة
    status_idx = pct_idx = None
    for idx, cell in enumerate(ws[1], 1):
        if str(cell.value) == "الحالة":
            status_idx = idx
        if str(cell.value) == "نسبة المطابقة %":
            pct_idx = idx

    for row in ws.iter_rows(min_row=2):
        row_fill = None
        if status_idx:
            sv       = str(row[status_idx - 1].value or "")
            row_fill = green_fill if "تمت المطابقة" in sv else red_fill
        for cell in row:
            if row_fill:
                cell.fill = row_fill
            cell.font      = normal_font
            cell.alignment = Alignment(horizontal="right", vertical="center")

        # تلوين خلية النسبة
        if pct_idx:
            pc   = row[pct_idx - 1]
            pstr = str(pc.value or "0").replace("%", "").strip()
            try:
                pv = float(pstr)
                if pv >= 95:
                    pc.fill = PatternFill("solid", fgColor="00B050")
                    pc.font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
                elif pv >= 80:
                    pc.fill = orange_fill
                else:
                    pc.fill = red_fill
            except ValueError:
                pass

    # عرض الأعمدة تلقائي
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[col_letter].width = min(max_len + 4, 55)

    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ══════════════════════════════════════════════════
#  مساعد: تخمين اسم العمود
# ══════════════════════════════════════════════════

def guess_col(cols: list, role: str) -> str | None:
    hints = {
        "payroll_name":   ["اسم", "name", "الاسم", "موظف"],
        "payroll_school": ["مدرسة", "قسم", "school", "المدرسة"],
        "db_name":        ["الاسم", "اسم", "name",  "موظف"],
        "db_school":      ["القسم", "قسم", "مدرسة", "school"],
    }
    for h in hints.get(role, []):
        for c in cols:
            if h in str(c).lower():
                return c
    return None


# ══════════════════════════════════════════════════
#  الواجهة الرئيسية
# ══════════════════════════════════════════════════

def main():
    # ── رأس الصفحة ──
    st.markdown("""
    <div class="page-header">
        <span>⚡</span>
        <div>
            <h1>تطبيق مطابقة الرواتب</h1>
            <p style="color:#8b949e;margin:0;font-size:.9rem">
                النسخة المطورة ٢.٠ — Streamlit
            </p>
        </div>
    </div>
    """, unsafe_allow_html=True)

    # ══════════════════════════════════════════════
    #  الشريط الجانبي — الإعدادات
    # ══════════════════════════════════════════════
    with st.sidebar:
        st.markdown("## 📂 تحميل الملفات")

        payroll_file = st.file_uploader(
            "ملف الرواتب / الأجور",
            type=["xlsx", "xls"],
            key="payroll_upload",
            help="Excel يحتوي على أسماء الموظفين ومدارسهم"
        )
        db_file = st.file_uploader(
            "قاعدة البيانات",
            type=["xlsx", "xls"],
            key="db_upload",
            help="Excel قاعدة بيانات الموظفين الكاملة"
        )

        st.divider()

        # ── قراءة الأعمدة ──
        payroll_df  = None
        payroll_cols: list = []
        db_cols:     list = []

        if payroll_file:
            try:
                payroll_df   = pd.read_excel(payroll_file, engine="openpyxl")
                payroll_cols = list(payroll_df.columns)
                st.success(f"✓ ملف الرواتب: **{len(payroll_df)}** صف")
            except Exception as e:
                st.error(f"فشل تحميل ملف الرواتب: {e}")

        if db_file:
            try:
                db_header = pd.read_excel(db_file, nrows=0, engine="openpyxl")
                db_cols   = list(db_header.columns)
                st.success(f"✓ قاعدة البيانات: **{len(db_cols)}** عمود")
            except Exception as e:
                st.error(f"فشل تحميل قاعدة البيانات: {e}")

        st.markdown("## 🗂 تعيين الأعمدة")

        p_name_def   = guess_col(payroll_cols, "payroll_name")
        p_school_def = guess_col(payroll_cols, "payroll_school")
        d_name_def   = guess_col(db_cols,      "db_name")
        d_school_def = guess_col(db_cols,      "db_school")

        payroll_name_col = st.selectbox(
            "اسم الموظف (ملف الأجور)",
            options=payroll_cols,
            index=payroll_cols.index(p_name_def) if p_name_def in payroll_cols else 0,
            disabled=not payroll_cols,
        )
        payroll_school_col = st.selectbox(
            "المدرسة / القسم (ملف الأجور)",
            options=["— بدون —"] + payroll_cols,
            index=(payroll_cols.index(p_school_def) + 1)
                  if p_school_def in payroll_cols else 0,
            disabled=not payroll_cols,
        )
        db_name_col = st.selectbox(
            "اسم الموظف (قاعدة البيانات)",
            options=db_cols,
            index=db_cols.index(d_name_def) if d_name_def in db_cols else 0,
            disabled=not db_cols,
        )
        db_school_col = st.selectbox(
            "المدرسة / القسم (قاعدة البيانات)",
            options=["— بدون —"] + db_cols,
            index=(db_cols.index(d_school_def) + 1)
                  if d_school_def in db_cols else 0,
            disabled=not db_cols,
        )

        # تحويل "— بدون —" → ""
        payroll_school_col = "" if payroll_school_col == "— بدون —" else payroll_school_col
        db_school_col      = "" if db_school_col      == "— بدون —" else db_school_col

        st.divider()
        st.markdown("## ⚙ إعدادات المطابقة")

        use_school = st.toggle("استخدام المدرسة / القسم", value=True)

        school_weight = st.slider(
            "وزن المدرسة", min_value=0, max_value=50,
            value=20, step=1,
            disabled=not use_school,
            help="كلما زاد الوزن زادت أهمية تطابق المدرسة"
        )
        threshold = st.slider(
            "عتبة المطابقة", min_value=80, max_value=130,
            value=100, step=1,
            help="≤95 تلتقط الفروق الإملائية البسيطة"
        )

        st.divider()
        st.markdown("## ➕ أعمدة إضافية من القاعدة")

        extra_cols: list = []
        if db_cols:
            extra_cols = st.multiselect(
                "اختر الأعمدة الإضافية",
                options=db_cols,
                default=[],
                help="ستُضاف هذه الأعمدة لكل صف في النتائج"
            )
        else:
            st.caption("حمّل قاعدة البيانات أولاً")

        st.divider()

        # ── زر التشغيل ──
        run_disabled = not (payroll_df is not None and db_file and
                            payroll_name_col and db_name_col)
        run_btn = st.button(
            "▶  بدء المطابقة",
            disabled=run_disabled,
            type="primary",
            use_container_width=True,
        )
        if run_disabled:
            st.caption("⚠ حمّل الملفين وحدد أعمدة الاسم أولاً")

    # ══════════════════════════════════════════════
    #  منطقة النتائج
    # ══════════════════════════════════════════════

    if run_btn and payroll_df is not None and db_file:
        # ── إعادة قراءة قاعدة البيانات بالأعمدة المطلوبة ──
        needed = list({db_name_col, db_school_col} | set(extra_cols))
        needed = [c for c in needed if c]          # حذف الفراغات

        try:
            db_file.seek(0)
            db_df = pd.read_excel(db_file, usecols=needed, engine="openpyxl")
        except Exception as e:
            st.error(f"خطأ في قراءة قاعدة البيانات: {e}")
            st.stop()

        records  = db_df.to_dict("records")
        db_index = build_index(records, db_name_col)

        payroll_rows = payroll_df.to_dict("records")
        total        = len(payroll_rows)

        # ── مؤشرات التقدم ──
        st.markdown("---")
        prog_bar    = st.progress(0, text="جاري المطابقة...")
        status_ph   = st.empty()

        # ── حاويات الإحصاء ──
        c1, c2, c3, c4 = st.columns(4)
        stat_matched = c1.empty()
        stat_missing = c2.empty()
        stat_review  = c3.empty()
        stat_total   = c4.empty()

        def render_stats(matched, missing, review, done):
            stat_matched.markdown(f"""
            <div class="stat-card">
                <div class="stat-lbl">✅ مطابق</div>
                <div class="stat-num" style="color:#3fb950">{matched}</div>
            </div>""", unsafe_allow_html=True)
            stat_missing.markdown(f"""
            <div class="stat-card">
                <div class="stat-lbl">❌ تدقيق بشري</div>
                <div class="stat-num" style="color:#da3633">{missing}</div>
            </div>""", unsafe_allow_html=True)
            stat_review.markdown(f"""
            <div class="stat-card">
                <div class="stat-lbl">⚠ مراجعة</div>
                <div class="stat-num" style="color:#d29922">{review}</div>
            </div>""", unsafe_allow_html=True)
            stat_total.markdown(f"""
            <div class="stat-card">
                <div class="stat-lbl">📋 المجموع</div>
                <div class="stat-num" style="color:#8b949e">{done}</div>
            </div>""", unsafe_allow_html=True)

        render_stats(0, 0, 0, 0)

        # ── المطابقة ──
        results      = []
        n_matched = n_missing = n_review = 0

        for i, row in enumerate(payroll_rows):
            name_req   = str(row.get(payroll_name_col, "") or "").strip()
            school_req = (str(row.get(payroll_school_col, "") or "").strip()
                          if payroll_school_col else "")

            best, final_score, name_score, match_pct = find_best_match(
                name_req, school_req,
                db_index, db_name_col, db_school_col,
                use_school, threshold, school_weight,
            )

            if best:
                matched_name   = best.get(db_name_col,   "")
                matched_school = best.get(db_school_col, "") if db_school_col else ""
                status_txt     = "✅ تمت المطابقة"
                n_matched     += 1
            else:
                matched_name   = "لم يتم العثور"
                matched_school = ""
                final_score = name_score = match_pct = 0.0
                status_txt  = "❌ تدقيق بشري"
                n_missing  += 1

            result_row: dict = {
                "#":                  i + 1,
                "اسم ملف الرواتب":    name_req,
                "المدرسة (الرواتب)":  school_req,
                "الاسم المطابق":      matched_name,
                "مدرسة القاعدة":      matched_school,
                "نسبة المطابقة %":    match_pct,
                "نقاط الاسم":         round(name_score,  1),
                "النقاط النهائية":    round(final_score, 1),
                "الحالة":             status_txt,
            }
            for col in extra_cols:
                result_row[col] = best.get(col, "") if best else ""

            results.append(result_row)

            # تحديث التقدم كل 30 صف
            if i % 30 == 0 or i == total - 1:
                pct = int((i + 1) / total * 100)
                prog_bar.progress(pct, text=f"جاري المطابقة... {i+1}/{total}")
                render_stats(n_matched, n_missing, n_review, i + 1)

        prog_bar.progress(100, text="✓ اكتملت المطابقة")
        status_ph.success(
            f"✓ اكتملت المطابقة — {n_matched} مطابق، {n_missing} للمراجعة"
        )

        # ══════════════════════════════════════════
        #  عرض النتائج
        # ══════════════════════════════════════════
        result_df = pd.DataFrame(results)

        st.markdown("---")
        st.markdown("### 📋 نتائج المطابقة")

        # ── فلتر الحالة ──
        col_f1, col_f2, col_f3 = st.columns([2, 2, 4])
        with col_f1:
            filter_status = st.selectbox(
                "فلتر الحالة",
                ["الكل", "✅ تمت المطابقة", "❌ تدقيق بشري"],
                key="filter_status",
            )
        with col_f2:
            filter_pct = st.slider(
                "نسبة المطابقة ≥", 80, 100, 0, key="filter_pct"
            )

        display_df = result_df.copy()
        if filter_status != "الكل":
            display_df = display_df[display_df["الحالة"] == filter_status]
        if filter_pct > 0:
            display_df = display_df[display_df["نسبة المطابقة %"] >= filter_pct]

        # ── تلوين الجدول ──
        def color_row(row):
            if "تمت المطابقة" in str(row["الحالة"]):
                return ["background-color:#0d2818; color:#e6edf3"] * len(row)
            return ["background-color:#2d0f0f; color:#e6edf3"] * len(row)

        def color_pct(val):
            try:
                v = float(val)
                if v >= 95:
                    return "background-color:#00B050; color:white; font-weight:bold"
                elif v >= 80:
                    return "background-color:#FFEB9C; color:#333"
                elif v > 0:
                    return "background-color:#FFC7CE; color:#333"
            except (TypeError, ValueError):
                pass
            return ""

        styled = (
            display_df.style
            .apply(color_row, axis=1)
            .map(color_pct, subset=["نسبة المطابقة %"])
            .format({"نسبة المطابقة %": "{:.0f}%"})
        )

        st.dataframe(
            styled,
            use_container_width=True,
            height=520,
        )

        st.caption(
            f"إجمالي المعروض: **{len(display_df)}** من أصل **{len(result_df)}** سجل"
        )

        # ══════════════════════════════════════════
        #  تنزيل Excel
        # ══════════════════════════════════════════
        st.markdown("---")

        # تحضير DataFrame للتصدير (النسبة كنص %)
        export_df = result_df.copy()
        export_df["نسبة المطابقة %"] = export_df["نسبة المطابقة %"].apply(
            lambda v: f"{v:.0f}%"
        )

        excel_bytes = build_excel(export_df)
        filename    = f"نتائج_المطابقة_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"

        col_dl1, col_dl2, _ = st.columns([2, 2, 4])
        with col_dl1:
            st.download_button(
                label="💾  تنزيل النتائج — Excel",
                data=excel_bytes,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
                use_container_width=True,
            )
        with col_dl2:
            csv_bytes = result_df.to_csv(index=False).encode("utf-8-sig")
            st.download_button(
                label="📄  تنزيل النتائج — CSV",
                data=csv_bytes,
                file_name=filename.replace(".xlsx", ".csv"),
                mime="text/csv",
                use_container_width=True,
            )

    else:
        # ── شاشة الترحيب ──
        st.markdown("""
        <div style="
            background:#161b22;
            border:1px solid #30363d;
            border-radius:12px;
            padding:48px 32px;
            text-align:center;
            margin-top:32px;
        ">
            <div style="font-size:3.5rem;margin-bottom:16px">📊</div>
            <h2 style="color:#e6edf3;margin:0 0 10px">جاهز للمطابقة</h2>
            <p style="color:#8b949e;font-size:1rem;max-width:440px;margin:0 auto">
                حمّل ملف الرواتب وقاعدة البيانات من الشريط الجانبي،
                ثم اضغط <b style="color:#3fb950">▶ بدء المطابقة</b>
            </p>
            <div style="
                display:flex;gap:20px;justify-content:center;
                margin-top:32px;flex-wrap:wrap
            ">
                <div style="background:#21262d;border-radius:8px;padding:14px 22px;min-width:160px">
                    <div style="font-size:1.6rem">🔤</div>
                    <div style="color:#c9d1d9;margin-top:6px;font-size:.9rem">
                        تطبيع عربي كامل<br>
                        <small style="color:#8b949e">أ إ آ ة ى ؤ ئ</small>
                    </div>
                </div>
                <div style="background:#21262d;border-radius:8px;padding:14px 22px;min-width:160px">
                    <div style="font-size:1.6rem">⚡</div>
                    <div style="color:#c9d1d9;margin-top:6px;font-size:.9rem">
                        فهرس ثنائي سريع<br>
                        <small style="color:#8b949e">آلاف السجلات في ثوانٍ</small>
                    </div>
                </div>
                <div style="background:#21262d;border-radius:8px;padding:14px 22px;min-width:160px">
                    <div style="font-size:1.6rem">🏫</div>
                    <div style="color:#c9d1d9;margin-top:6px;font-size:.9rem">
                        دعم المدرسة / القسم<br>
                        <small style="color:#8b949e">لتحسين دقة المطابقة</small>
                    </div>
                </div>
                <div style="background:#21262d;border-radius:8px;padding:14px 22px;min-width:160px">
                    <div style="font-size:1.6rem">📊</div>
                    <div style="color:#c9d1d9;margin-top:6px;font-size:.9rem">
                        نسبة المطابقة %<br>
                        <small style="color:#8b949e">تلوين تلقائي بالجودة</small>
                    </div>
                </div>
            </div>
        </div>
        """, unsafe_allow_html=True)


if __name__ == "__main__":
    main()
